"""Bounded attachment extraction in a killable process. Never executes files."""

import io
import multiprocessing
from pathlib import PurePath
import zipfile

from .console_models import ConsoleError

MAX_BYTES = 20*1024*1024
MAX_TEXT = 12000


def _zip_check(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(i.file_size for i in entries) > 80*1024*1024:
            raise ConsoleError("archive_expansion_limit")
        if any(i.file_size > 20*1024*1024 or i.flag_bits & 1 for i in entries):
            raise ConsoleError("archive_member_limit_or_encrypted")


def extract_bytes(name: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ConsoleError("attachment_too_large")
    suffix = PurePath(name).suffix.lower()
    if suffix in (".txt", ".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeError:
            try:
                text = data.decode("gb18030")
            except UnicodeError:
                raise ConsoleError("unsupported_text_encoding") from None
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ConsoleError("encrypted_document")
        pieces = []
        for page in reader.pages[:100]:
            pieces.append(page.extract_text() or "")
            if sum(map(len, pieces)) >= MAX_TEXT:
                break
        text = "\n".join(pieces)
    elif suffix == ".docx":
        _zip_check(data)
        from docx import Document
        document = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n"+" | ".join(cell.text for cell in row.cells)
                if len(text) >= MAX_TEXT:
                    break
            if len(text) >= MAX_TEXT:
                break
    elif suffix == ".xlsx":
        _zip_check(data)
        import defusedxml  # required before openpyxl; no XML-entity fallback
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        pieces = []
        try:
            for sheet in workbook.worksheets[:20]:
                pieces.append("工作表："+sheet.title)
                for row in sheet.iter_rows(max_row=1000, max_col=50, values_only=True):
                    pieces.append(" | ".join("" if cell is None else str(cell)[:500] for cell in row))
                    if sum(map(len, pieces)) >= MAX_TEXT:
                        break
                if sum(map(len, pieces)) >= MAX_TEXT:
                    break
            text = "\n".join(pieces)
        finally:
            workbook.close()
    else:
        raise ConsoleError("unsupported_attachment_type")
    if not text.strip():
        raise ConsoleError("no_extractable_text")
    return text[:MAX_TEXT] + ("\n[内容已截断]" if len(text) > MAX_TEXT else "")


def _extract_worker(connection, name, data):
    try:
        connection.send({"text": extract_bytes(name, data)})
    except ConsoleError as exc:
        connection.send({"error": str(exc)})
    except ImportError:
        connection.send({"error": "document_dependency_missing"})
    except Exception:
        connection.send({"error": "document_parse_failed"})
    finally:
        connection.close()


def extract(name: str, data: bytes, timeout=30) -> str:
    if len(data) > MAX_BYTES:
        raise ConsoleError("attachment_too_large")
    ctx = multiprocessing.get_context("spawn")
    parent, child = ctx.Pipe(duplex=False)
    process = ctx.Process(target=_extract_worker, args=(child, name, data), daemon=True)
    process.start()
    child.close()
    try:
        if not parent.poll(timeout):
            raise ConsoleError("document_parse_timeout")
        try:
            result = parent.recv()
        except EOFError:
            raise ConsoleError("document_parse_failed") from None
        if "error" in result:
            raise ConsoleError(result["error"])
        return result["text"]
    finally:
        if process.is_alive():
            process.terminate()
        process.join(timeout=2)
        parent.close()
